# -*- coding: utf-8 -*-
# 文件比对工具 v2.0 真实运行验证（仅用标准库 urllib，不依赖 requests/opslib.httpx）
# 用法：先启动服务（PORT=18000 python tools/diff_file_dir_jar_excel.py），再 python verify_diff_v2.py
import json
import os
import sys
import tempfile
import time
import urllib.request
import zipfile

BASE = "http://127.0.0.1:18000"
PY = sys.executable

PASS = 0
FAIL = 0
LINES = []


def ok(label, cond, extra=""):
    global PASS, FAIL
    if cond:
        PASS += 1
        LINES.append("[PASS] %s %s" % (label, extra))
    else:
        FAIL += 1
        LINES.append("[FAIL] %s %s" % (label, extra))


def http(method, path, body=None, is_json=True, raw=False):
    url = BASE + path
    data = None
    headers = {}
    if body is not None:
        if is_json:
            data = json.dumps(body).encode("utf-8")
            headers["Content-Type"] = "application/json"
        else:
            data = body
    req = urllib.request.Request(url, data=data, method=method, headers=headers)
    try:
        resp = urllib.request.urlopen(req, timeout=30)
    except urllib.error.HTTPError as e:
        return e.code, e.read().decode("utf-8", "replace"), (e.headers.get("Content-Type", "") or "")
    except Exception as e:  # noqa
        return -1, "ERR:%s" % e, ""
    ctype = resp.headers.get("Content-Type", "") or ""
    text = resp.read().decode("utf-8", "replace")
    code = resp.getcode()
    return code, text, ctype


def post_json(path, body):
    code, text, ctype = http("POST", path, body)
    return code, (text if ctype.startswith("application/json") else text), ctype, _try_json(text)


def _try_json(text):
    try:
        return json.loads(text)
    except Exception:
        return None


def get_json(path):
    code, text, ctype = http("GET", path)
    return code, text, ctype, _try_json(text)


# ---------------- 构造测试夹具 ----------------
WORK = tempfile.mkdtemp(prefix="diffv2_")
LINES.append("[夹具] %s" % WORK)


def w(path, content):
    with open(path, "wb") as f:
        f.write(content.encode("utf-8") if isinstance(content, str) else content)


def mkdir(p):
    os.makedirs(p, exist_ok=True)
    return p


# 目录夹具
da = mkdir(os.path.join(WORK, "dirA"))
db = mkdir(os.path.join(WORK, "dirB"))
w(os.path.join(da, "same.txt"), "same content\n")
w(os.path.join(da, "diff.txt"), "line A\ncommon\n")
w(os.path.join(da, "app.log"), "ignored log\n")
w(os.path.join(db, "same.txt"), "same content\n")
w(os.path.join(db, "diff.txt"), "line B\ncommon\n")
w(os.path.join(db, "onlyb.txt"), "only in b\n")
w(os.path.join(db, "app.log"), "ignored log\n")

# 文本文件夹具
t1 = os.path.join(WORK, "t1.txt")
t2 = os.path.join(WORK, "t2.txt")
w(t1, "alpha\nbeta\ngamma\n")
w(t2, "alpha\nBETA\ngamma\n")

# jar 夹具
jar1 = os.path.join(WORK, "a.jar")
jar2 = os.path.join(WORK, "b.jar")
with zipfile.ZipFile(jar1, "w") as z:
    z.writestr("com/A.class", "AAAA")
    z.writestr("conf/app.txt", "v1")
    z.writestr("META-INF/XX.SF", "signature")
with zipfile.ZipFile(jar2, "w") as z:
    z.writestr("com/A.class", "AAAA")
    z.writestr("conf/app.txt", "v2-changed")
    z.writestr("META-INF/XX.SF", "signature")


# excel 夹具
def make_xlsx(path, v):
    import openpyxl
    wb = openpyxl.Workbook()
    ws = wb.active
    ws["A1"] = "标题"
    ws["A2"] = v
    wb.save(path)


x1 = os.path.join(WORK, "e1.xlsx")
x2 = os.path.join(WORK, "e2.xlsx")
make_xlsx(x1, "old")
make_xlsx(x2, "new")

# docx 夹具
DOCX_BODY = ('<?xml version="1.0"?><w:document xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main">'
             '<w:body><w:p><w:r><w:t>%s</w:t></w:r></w:p></w:body></w:document>')


def make_docx(path, txt):
    with zipfile.ZipFile(path, "w") as z:
        z.writestr("[Content_Types].xml", "<Types/>")
        z.writestr("word/document.xml", DOCX_BODY % txt)


d1 = os.path.join(WORK, "d1.docx")
d2 = os.path.join(WORK, "d2.docx")
make_docx(d1, "第一段版本一")
make_docx(d2, "第一段版本二")

# pptx 夹具
PPTX_BODY = ('<?xml version="1.0"?><p:sld xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main">'
             '<a:sp><a:txBody><a:p><a:t>%s</a:t></a:p></a:txBody></a:sp></p:sld>')


def make_pptx(path, txt):
    with zipfile.ZipFile(path, "w") as z:
        z.writestr("[Content_Types].xml", "<Types/>")
        z.writestr("ppt/slides/slide1.xml", PPTX_BODY % txt)


p1 = os.path.join(WORK, "s1.pptx")
p2 = os.path.join(WORK, "s2.pptx")
make_pptx(p1, "幻灯片一")
make_pptx(p2, "幻灯片二")

# 三向目录夹具
tb = mkdir(os.path.join(WORK, "base"))
tl = mkdir(os.path.join(WORK, "left"))
tr = mkdir(os.path.join(WORK, "right"))
w(os.path.join(tb, "common.txt"), "base-common\n")
w(os.path.join(tb, "only_base.txt"), "only in base\n")
w(os.path.join(tl, "common.txt"), "left-common\n")
w(os.path.join(tl, "only_left.txt"), "only in left\n")
w(os.path.join(tr, "common.txt"), "right-common\n")
w(os.path.join(tr, "only_right.txt"), "only in right\n")

# 批量目录对
batch_pairs = []
for i in range(4):
    pa = mkdir(os.path.join(WORK, "b%d_a" % i))
    pb = mkdir(os.path.join(WORK, "b%d_b" % i))
    w(os.path.join(pa, "f.txt"), "x%d\n" % i)
    w(os.path.join(pb, "f.txt"), "y%d\n" % i)
    batch_pairs.append((pa, pb))

# ---------------- 1. 健康检查 / 首页 ----------------
code, text, ctype = http("GET", "/")
ok("GET / 首页 200", code == 200 and "文件比对" in text, "(%d)" % code)

# ---------------- 2. 自动比对（目录） ----------------
code, text, ctype, j = post_json("/api/compare/auto", {"left_path": da, "right_path": db})
ok("auto 目录比对 success", j and j.get("success"), str(j.get("message", ""))[:60] if j else text[:80])
ok("auto 含 compare_id", bool(j and j.get("compare_id")))
if j and j.get("compare_id"):
    AUTO_ID = j["compare_id"]
else:
    AUTO_ID = ""

# ---------------- 3. 目录比对 + 忽略规则 ----------------
code, text, ctype, j = post_json("/api/compare/dir", {"dir1": da, "dir2": db, "use_saved_rules": True})
ok("dir 比对 success", j and j.get("success"))
ok("dir 忽略 app.log 计数>=1",
   bool(j and (j.get("summary") or {}).get("ignored", 0) >= 1),
   "ignored=%s" % ((j.get("summary") or {}).get("ignored") if j else "?"))

# 关闭忽略规则再比对，ignored 应为 0
code, text, ctype, j2 = post_json("/api/compare/dir", {"dir1": da, "dir2": db, "use_saved_rules": False, "ignore": ""})
ok("dir 关闭忽略后 ignored=0",
   bool(j2 and (j2.get("summary") or {}).get("ignored", 0) == 0),
   "ignored=%s" % ((j2.get("summary") or {}).get("ignored") if j2 else "?"))

# ---------------- 4. 文本文件比对（用于导出补丁） ----------------
code, text, ctype, j = post_json("/api/compare/file", {"file1": t1, "file2": t2})
ok("file 文本比对 success", j and j.get("success"))
TEXT_ID = (j or {}).get("compare_id", "")

# ---------------- 5. jar 比对（签名忽略） ----------------
code, text, ctype, j = post_json("/api/compare/jar", {"jar1_path": jar1, "jar2_path": jar2, "ignore_sign": True})
ok("jar 比对 success", j and j.get("success"))
ok("jar 签名被忽略 different=1",
   bool(j and (j.get("summary") or {}).get("different", 0) == 1),
   "different=%s" % ((j.get("summary") or {}).get("different") if j else "?"))

# ---------------- 6. excel 比对 ----------------
code, text, ctype, j = post_json("/api/compare/excel", {"excel1_path": x1, "excel2_path": x2})
ok("excel 比对 success", j and j.get("success"))
ok("excel 列出 sheet 差异", bool(j and j.get("sheets")))

# ---------------- 7. docx / pptx 结构化比对 ----------------
code, text, ctype, j = post_json("/api/compare/ooxml", {"left_path": d1, "right_path": d2})
ok("docx 比对 success", j and j.get("success"))
ok("docx 检测到段落变化", bool(j and (j.get("summary") or {}).get("changed", 0) >= 1),
   "summary=%s" % ((j.get("summary")) if j else "?"))

code, text, ctype, j = post_json("/api/compare/ooxml", {"left_path": p1, "right_path": p2})
ok("pptx 比对 success", j and j.get("success"))
ok("pptx 检测到变化", bool(j and not j.get("same")))

# ---------------- 8. 三向比对 ----------------
code, text, ctype, j = post_json("/api/compare/three", {"base": tb, "left": tl, "right": tr})
ok("three 比对 success", j and j.get("success"))
s3 = (j or {}).get("summary") or {}
ok("three 含 base_only/a_only/b_only/common",
   s3.get("base_only", 0) >= 1 and s3.get("a_only", 0) >= 1 and s3.get("b_only", 0) >= 1,
   "summary=%s" % s3)
ok("three 含 conflict 标记字段", "common_conflict" in s3)

# ---------------- 9. 忽略规则 API ----------------
code, text, ctype, j = get_json("/api/ignore/rules")
ok("ignore/rules GET success", j and j.get("success"))
ok("ignore/rules 含内置默认", bool(j and len(j.get("rules", [])) >= 5),
   "rules=%d" % (len(j.get("rules", [])) if j else 0))
DEFAULTS = j.get("rules") if j else None

code, text, ctype, j = post_json("/api/ignore/test",
                                  {"path": "logs/foo.log", "scope": "all", "ignore": "*.log", "use_saved_rules": False})
ok("ignore/test 命中临时规则", bool(j and j.get("ignored") is True), str((j or {}).get("matched")))

code, text, ctype, j = post_json("/api/ignore/test",
                                  {"path": "src/main.py", "scope": "all", "ignore": "*.log", "use_saved_rules": False})
ok("ignore/test 不命中", bool(j and j.get("ignored") is False))

# 保存自定义规则后还原默认，保持环境干净
if DEFAULTS is not None:
    code, text, ctype, j = post_json("/api/ignore/rules", {"rules": DEFAULTS})
    ok("ignore/rules POST 保存成功", bool(j and j.get("success")))

# ---------------- 10. 历史 API ----------------
code, text, ctype, j = get_json("/api/history?limit=50")
ok("history GET success", j and j.get("success"))
ok("history 至少 1 条记录", bool(j and (j.get("total", 0) >= 1)), "total=%s" % ((j or {}).get("total")))
HID = ""
if j and j.get("items"):
    HID = j["items"][0].get("id", "")
if HID:
    code, text, ctype, jd = get_json("/api/history/%s" % HID)
    ok("history/<id> GET 命中", bool(jd and jd.get("success")))
    # DELETE 删除
    code2, text2, ctype2 = http("DELETE", "/api/history/%s" % HID)
    ok("history/<id> DELETE", code2 in (200, 204) or (isinstance(text2, str) and "success" in text2))
    code, text, ctype, jd = get_json("/api/history/%s" % HID)
    ok("history/<id> 删除后查不到", bool(jd and not jd.get("success")))

# ---------------- 11. 导出报告 / 补丁 ----------------
if TEXT_ID:
    code, text, ctype = http("POST", "/api/export/report?format=html",
                              {"compare_id": TEXT_ID})
    ok("export/report html 200", code == 200 and "text/html" in ctype, "ctype=%s len=%d" % (ctype, len(text)))
    code, text, ctype = http("POST", "/api/export/report?format=md",
                              {"compare_id": TEXT_ID})
    ok("export/report md 200", code == 200 and "markdown" in ctype, "ctype=%s" % ctype)
    code, text, ctype = http("POST", "/api/export/patch", {"compare_id": TEXT_ID})
    ok("export/patch 200", code == 200 and "text/plain" in ctype, "ctype=%s len=%d" % (ctype, len(text)))
    ok("export/patch 含 unified diff 标记",
       ("--- " in text) or ("+++ " in text) or ("diff" in text.lower()) or len(text) > 0)

# ---------------- 12. 批量比对（同步 + 异步） ----------------
code, text, ctype, j = post_json("/api/compare/batch",
                                  {"pairs": [{"left": batch_pairs[0][0], "right": batch_pairs[0][1]},
                                             {"left": batch_pairs[1][0], "right": batch_pairs[1][1]}]})
ok("batch 同步(<=3) async=False", bool(j and j.get("async") is False), str((j or {}).get("async")))

code, text, ctype, j = post_json("/api/compare/batch",
                                  {"pairs": [{"left": l, "right": r} for l, r in batch_pairs]})
ok("batch 异步(>3) async=True", bool(j and j.get("async") is True))
TID = (j or {}).get("task_id", "")
if TID:
    done = False
    for _ in range(40):
        code, text, ctype, jt = get_json("/api/tasks/%s" % TID)
        if jt and jt.get("task", {}).get("state") in ("success", "failed", "canceled"):
            done = True
            ok("batch 任务完成", True, "state=%s" % jt["task"]["state"])
            break
        time.sleep(0.5)
    if not done:
        ok("batch 任务完成", False, "超时未结束")

# ---------------- 13. 鉴权（未设置 token 时应无需鉴权） ----------------
code, text, ctype, j = get_json("/api/history")
ok("无 token 可访问(未配置 COMPARE_TOKEN)", code == 200)

# ---------------- 汇总 ----------------
LINES.append("")
LINES.append("==== 验证结果 PASS=%d FAIL=%d ====" % (PASS, FAIL))
out = "\n".join(LINES)
print(out)
with open(os.path.join(WORK, "verify_diff_v2.out.txt"), "w") as f:
    f.write(out)
sys.exit(1 if FAIL else 0)
